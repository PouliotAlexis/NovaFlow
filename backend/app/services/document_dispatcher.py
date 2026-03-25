import os
import logging
import fitz  # PyMuPDF
from docx import Document
from pptx import Presentation
from typing import Optional

logger = logging.getLogger("document_dispatcher")

class DocumentDispatcher:
    def __init__(self):
        pass

    def extract_text(self, file_path: str) -> Optional[str]:
        """Extrait le texte d'un fichier selon son extension."""
        ext = os.path.splitext(file_path)[1].lower()
        
        try:
            if ext == ".pdf":
                return self._extract_pdf(file_path)
            elif ext == ".docx":
                return self._extract_docx(file_path)
            elif ext in [".pptx", ".ppt"]:
                return self._extract_pptx(file_path)
            elif ext in [".txt", ".md"]:
                with open(file_path, "r", encoding="utf-8") as f:
                    return f.read()
            else:
                logger.warning(f"Format non supporté pour l'extraction: {ext}")
                return None
        except Exception as e:
            logger.error(f"Erreur d'extraction pour {file_path}: {e}")
            return None

    def _extract_pdf(self, file_path: str) -> str:
        text = ""
        with fitz.open(file_path) as doc:
            for page in doc:
                text += page.get_text()
        return text

    def _extract_docx(self, file_path: str) -> str:
        doc = Document(file_path)
        return "\n".join([para.text for para in doc.paragraphs])

    def _extract_pptx(self, file_path: str) -> str:
        prs = Presentation(file_path)
        text_runs = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text_runs.append(shape.text)
        return "\n".join(text_runs)

    def normalize_to_markdown(self, text: str, title: str) -> str:
        """Enveloppe le texte dans une structure Markdown simple."""
        return f"# {title}\n\n{text}"

dispatcher = DocumentDispatcher()
